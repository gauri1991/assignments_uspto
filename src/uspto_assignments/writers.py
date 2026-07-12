"""Streaming file writers for the full record stream.

Both writers stream so peak memory stays flat on multi-GB inputs: Parquet flushes in batches
via :class:`pyarrow.parquet.ParquetWriter`, and Excel uses openpyxl ``write_only`` mode (one
temp file per sheet). These consume the raw :func:`~uspto_assignments.parser.iter_records`
stream; the in-memory/interactive path (Phase 2) reuses the same table transforms.
"""

from __future__ import annotations

import logging
from collections.abc import Iterator
from dataclasses import fields
from pathlib import Path
from typing import Any

import pyarrow as pa
import pyarrow.parquet as pq

from .model import (
    DEFAULT_BATCH_SIZE,
    EXCEL_MAX_ROWS,
    FLAT_COLUMNS,
    TABLE_TYPES,
    ExtractedRecord,
    schema_for,
)
from .tables import flat_rows, rows_to_table

__all__ = ["DEFAULT_BATCH_SIZE", "EXCEL_MAX_ROWS", "write_excel", "write_parquet"]

logger = logging.getLogger(__name__)


class _ParquetSink:
    """Batched Parquet writer for one table; opens the file lazily on first flush."""

    def __init__(self, out_path: Path, row_type: type) -> None:
        self._path = out_path
        self._row_type = row_type
        self._schema = schema_for(row_type)
        self._writer: pq.ParquetWriter | None = None
        self.row_count = 0

    def write(self, rows: list[Any]) -> None:
        if not rows:
            return
        if self._writer is None:
            self._writer = pq.ParquetWriter(str(self._path), self._schema)
        self._writer.write_table(rows_to_table(rows, self._row_type))
        self.row_count += len(rows)

    def close(self) -> None:
        if self._writer is None:
            # Nothing was written — still emit an empty, well-typed file for consistency.
            pq.write_table(  # pyright: ignore[reportUnknownMemberType]  # stub embeds Unknown
                pa.table({f.name: [] for f in fields(self._row_type)}, self._schema),
                str(self._path),
            )
        else:
            self._writer.close()


def write_parquet(
    records: Iterator[ExtractedRecord],
    outdir: Path,
    basename: str,
    batch_size: int,
) -> dict[str, int]:
    """Stream records to one Parquet file per table, flushing every ``batch_size`` records.

    Args:
        records: Stream of extracted records.
        outdir: Directory to write ``<table>.parquet`` files into.
        basename: Unused for Parquet naming (kept for symmetry with the Excel writer).
        batch_size: Number of assignment records to buffer before each flush.

    Returns:
        Mapping of table name to the number of rows written.
    """
    sinks = {name: _ParquetSink(outdir / f"{name}.parquet", t) for name, t in TABLE_TYPES.items()}
    buffers: dict[str, list[Any]] = {name: [] for name in TABLE_TYPES}
    seen = 0

    def flush() -> None:
        for name, sink in sinks.items():
            sink.write(buffers[name])
            buffers[name].clear()

    for rec in records:
        buffers["assignments"].append(rec.assignment)
        buffers["assignors"].extend(rec.assignors)
        buffers["assignees"].extend(rec.assignees)
        buffers["properties"].extend(rec.properties)
        seen += 1
        if seen % batch_size == 0:
            flush()
            logger.info("Parsed %d assignments...", seen)
    flush()
    for sink in sinks.values():
        sink.close()
    logger.info("Finished: %d assignments parsed", seen)
    return {name: sink.row_count for name, sink in sinks.items()}


def write_excel(
    records: Iterator[ExtractedRecord],
    outdir: Path,
    basename: str,
) -> dict[str, int]:
    """Write all tables (+ flat view) to a single multi-sheet Excel workbook.

    Excel caps a sheet at ``EXCEL_MAX_ROWS`` rows; any table exceeding that is truncated on
    the sheet (with a warning) — the Parquet output remains the complete source of truth.

    Args:
        records: Stream of extracted records.
        outdir: Directory to write ``<basename>.xlsx`` into.
        basename: Workbook filename stem.

    Returns:
        Mapping of sheet name to rows written to that sheet (post-truncation).
    """
    # Imported lazily so a Parquet-only run skips openpyxl's import cost. write_only mode
    # streams rows to a temp file per sheet, so peak memory stays flat on huge inputs.
    from openpyxl import Workbook  # noqa: PLC0415

    usable = EXCEL_MAX_ROWS - 1  # reserve one row for the header
    table_columns = {name: [f.name for f in fields(t)] for name, t in TABLE_TYPES.items()}
    sheet_columns: dict[str, list[str]] = {**table_columns, "flat": FLAT_COLUMNS}

    wb = Workbook(write_only=True)
    sheets = {name: wb.create_sheet(title=name) for name in sheet_columns}
    for name, cols in sheet_columns.items():
        sheets[name].append(cols)

    counts = dict.fromkeys(sheet_columns, 0)
    truncated: set[str] = set()

    def emit(name: str, values: list[Any]) -> None:
        if counts[name] >= usable:
            truncated.add(name)
            return
        sheets[name].append(values)
        counts[name] += 1

    for rec in records:
        emit("assignments", [getattr(rec.assignment, c) for c in table_columns["assignments"]])
        for assignor in rec.assignors:
            emit("assignors", [getattr(assignor, c) for c in table_columns["assignors"]])
        for assignee in rec.assignees:
            emit("assignees", [getattr(assignee, c) for c in table_columns["assignees"]])
        for prop in rec.properties:
            emit("properties", [getattr(prop, c) for c in table_columns["properties"]])
        for flat_row in flat_rows(rec):
            emit("flat", [flat_row[c] for c in FLAT_COLUMNS])

    out_path = outdir / f"{basename}.xlsx"
    wb.save(str(out_path))
    for name in sorted(truncated):
        logger.warning(
            "Sheet %r hit Excel's row limit (%d) and was truncated; the Parquet %r is complete.",
            name,
            usable,
            name,
        )
    return counts
